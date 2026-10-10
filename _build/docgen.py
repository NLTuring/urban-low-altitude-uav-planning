# -*- coding: utf-8 -*-
"""Building blocks for the project brief DOCX. Font: SimSun (宋体) throughout."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_LINE_SPACING,
                            WD_TAB_ALIGNMENT, WD_TAB_LEADER)
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = '宋体'
BODY_SIZE = 12.0        # 小四
TABLE_SIZE = 10.5       # 五号
USABLE_CM = 15.4        # A4 21cm - 2.8cm * 2

GREY = 'F2F2F2'
GREY_HEAD = 'E8E8E8'
ACCENT = '1F4E79'


# --------------------------------------------------------------------------
# low level xml helpers
# --------------------------------------------------------------------------
def _set_rfonts(rpr, name):
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(attr), name)


def style_run(run, size=BODY_SIZE, bold=False, italic=False, color=None, font=FONT):
    run.font.name = font
    _set_rfonts(run._element.get_or_add_rPr(), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def para_shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), fill)
    pPr.append(sh)


def cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), fill)
    tcPr.append(sh)


def para_border(p, edges=('left',), sz=18, color=ACCENT):
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn('w:pBdr'))
    if bdr is None:
        bdr = OxmlElement('w:pBdr')
        pPr.append(bdr)
    for e in edges:
        el = OxmlElement('w:' + e)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '6')
        el.set(qn('w:color'), color)
        bdr.append(el)


def indent_chars(p, first_line_chars=None, left_chars=None, hanging_chars=None):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    if first_line_chars is not None:
        ind.set(qn('w:firstLineChars'), str(first_line_chars))
        ind.set(qn('w:firstLine'), str(int(first_line_chars / 100.0 * BODY_SIZE * 20)))
    if left_chars is not None:
        ind.set(qn('w:leftChars'), str(left_chars))
        ind.set(qn('w:left'), str(int(left_chars / 100.0 * BODY_SIZE * 20)))
    if hanging_chars is not None:
        ind.set(qn('w:hangingChars'), str(hanging_chars))
        ind.set(qn('w:hanging'), str(int(hanging_chars / 100.0 * BODY_SIZE * 20)))


def keep_with_next(p):
    # w:pPr children must follow the schema order (pStyle, keepNext, keepLines,
    # ...). Appending at the end produces a file LibreOffice silently ignores,
    # which is exactly how a table caption ends up orphaned at a page bottom.
    p.paragraph_format.keep_with_next = True


def para_mark_size(p, size_pt):
    """Set the paragraph-mark (¶) font size.

    The pilcrow inherits the Normal style size, and since it participates in
    line height it silently inflates densely packed paragraphs (a TOC list, for
    example) so a list that should fit one page spills a few lines onto the
    next. w:rPr is late in the pPr sequence, so appending is schema-correct.
    """
    pPr = p._p.get_or_add_pPr()
    rPr = pPr.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        pPr.append(rPr)
    _set_rfonts(rPr, FONT)
    for tag in ('w:sz', 'w:szCs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rPr.append(el)
        el.set(qn('w:val'), str(int(round(size_pt * 2))))


_TRPR_SEQ = (
    'w:cnfStyle', 'w:divId', 'w:gridBefore', 'w:gridAfter', 'w:wBefore',
    'w:wAfter', 'w:cantSplit', 'w:trHeight', 'w:tblHeader', 'w:tblCellSpacing',
    'w:jc', 'w:hidden', 'w:ins', 'w:del', 'w:trPrChange',
)


def _trpr_set(row, tag):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn(tag)) is not None:
        return
    el = OxmlElement(tag)
    idx = _TRPR_SEQ.index(tag)
    anchor = None
    for child in trPr:
        name = 'w:' + child.tag.split('}')[-1]
        if name in _TRPR_SEQ and _TRPR_SEQ.index(name) > idx:
            anchor = child
            break
    if anchor is None:
        trPr.append(el)
    else:
        anchor.addprevious(el)


def cant_split(row):
    _trpr_set(row, 'w:cantSplit')


def table_header_repeat(row):
    _trpr_set(row, 'w:tblHeader')


def add_field(paragraph, instr, placeholder=''):
    run = paragraph.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = placeholder
    f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2); run._r.append(t); run._r.append(f3)
    return run


# --------------------------------------------------------------------------
# document setup
# --------------------------------------------------------------------------
def new_document():
    doc = Document()

    # page setup
    for s in doc.sections:
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.top_margin = Cm(2.54)
        s.bottom_margin = Cm(2.54)
        s.left_margin = Cm(2.8)
        s.right_margin = Cm(2.8)
        s.header_distance = Cm(1.5)
        s.footer_distance = Cm(1.5)

    # base style
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(BODY_SIZE)
    _set_rfonts(normal.element.get_or_add_rPr(), FONT)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    heading_spec = {
        'Heading 1': (16, 18, 12),
        'Heading 2': (14, 14, 8),
        'Heading 3': (12, 10, 6),
    }
    for name, (size, before, after) in heading_spec.items():
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string('000000')
        _set_rfonts(st.element.get_or_add_rPr(), FONT)
        pf = st.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = 1.4
        pf.keep_with_next = True
        pf.page_break_before = False

    for name in ('TOC 1', 'TOC 2', 'TOC Heading'):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = FONT
        st.font.size = Pt(BODY_SIZE if name == 'TOC Heading' else TABLE_SIZE)
        st.font.bold = (name == 'TOC Heading')
        st.font.color.rgb = RGBColor.from_string('000000')
        _set_rfonts(st.element.get_or_add_rPr(), FONT)
        st.paragraph_format.line_spacing = 1.18
        st.paragraph_format.space_after = Pt(0)

    return doc


def repeat_header_row(row):
    table_header_repeat(row)
    for c in row.cells:
        cell_shade(c, GREY_HEAD)
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True


# --------------------------------------------------------------------------
# content blocks
# --------------------------------------------------------------------------
def h1(doc, text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        style_run(r, size=16, bold=True)
    return p


def h2(doc, text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        style_run(r, size=14, bold=True)
    return p


def h3(doc, text):
    p = doc.add_heading(text, level=3)
    for r in p.runs:
        style_run(r, size=12, bold=True)
    return p


def p(doc, text, indent=True):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.paragraph_format.space_after = Pt(3)
    run = par.add_run(text)
    style_run(run, size=BODY_SIZE)
    if indent:
        indent_chars(par, first_line_chars=200)
    return par


def bullet(doc, text, level=1):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.line_spacing = 1.4
    marker = '● ' if level == 1 else '○ '
    left = 100 if level == 1 else 200
    run = par.add_run(marker + text)
    style_run(run, size=BODY_SIZE)
    indent_chars(par, left_chars=left, hanging_chars=100)
    return par


def note(doc, text, fill=GREY, label=None):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.space_after = Pt(8)
    par.paragraph_format.line_spacing = 1.35
    if label:
        r0 = par.add_run(label)
        style_run(r0, size=TABLE_SIZE, bold=True)
    run = par.add_run(text)
    style_run(run, size=TABLE_SIZE)
    para_shade(par, fill)
    para_border(par, edges=('left',))
    indent_chars(par, left_chars=50)
    return par


def fig(doc, text):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.space_after = Pt(10)
    par.paragraph_format.line_spacing = 1.2
    run = par.add_run('【图示】' + text)
    style_run(run, size=TABLE_SIZE, italic=False, color='7F7F7F')
    para_border(par, edges=('top', 'bottom', 'left', 'right'), sz=6, color='BFBFBF')
    return par


def caption(doc, text):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(10)
    run = par.add_run(text)
    style_run(run, size=TABLE_SIZE, bold=True)
    # a table caption must never be left alone at the bottom of a page
    keep_with_next(par)
    return par


def table(doc, cols, rows, widths=None, caption_text=None, first_col_bold=False):
    if caption_text:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cap.paragraph_format.space_before = Pt(8)
        cap.paragraph_format.space_after = Pt(4)
        r = cap.add_run(caption_text)
        style_run(r, size=TABLE_SIZE, bold=True)
        keep_with_next(cap)

    t = doc.add_table(rows=1, cols=len(cols))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    hdr = t.rows[0]
    for i, c in enumerate(cols):
        cell = hdr.cells[i]
        cell.text = ''
        par = cell.paragraphs[0]
        par.paragraph_format.line_spacing = 1.15
        par.paragraph_format.space_after = Pt(0)
        run = par.add_run(c)
        style_run(run, size=TABLE_SIZE, bold=True)
    repeat_header_row(hdr)
    cant_split(hdr)
    # keep the caption, the header row and the first data row on one page: the
    # caption's keepNext alone does not survive a table boundary in every
    # renderer, so chain the first two rows into the same group.
    for _c in hdr.cells:
        for _p in _c.paragraphs:
            _p.paragraph_format.keep_with_next = True

    for row_data in rows:
        row = t.add_row()
        cant_split(row)
        first_data_row = (row is t.rows[1])
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.text = ''
            par = cell.paragraphs[0]
            par.paragraph_format.line_spacing = 1.2
            par.paragraph_format.space_after = Pt(0)
            if first_data_row:
                par.paragraph_format.keep_with_next = True
            bold = first_col_bold and i == 0
            run = par.add_run(str(val))
            style_run(run, size=TABLE_SIZE, bold=bold)

    if widths:
        total = sum(widths)
        for i, w in enumerate(widths):
            cm = USABLE_CM * w / total
            for row in t.rows:
                row.cells[i].width = Cm(cm)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.line_spacing = 1.0
    for r in sp.runs:
        style_run(r, size=6)
    return t


def page_break(doc):
    doc.add_page_break()


# --------------------------------------------------------------------------
# cover / toc / sections
# --------------------------------------------------------------------------
def cover(doc, titles, meta, banner_line='参 赛 作 品 项 目 简 介',
          subtitle='—— 商业计划书与路演 PPT 编制指南 ——'):
    def line(text, size, bold=False, before=0, after=0, spacing=1.3):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(before)
        par.paragraph_format.space_after = Pt(after)
        par.paragraph_format.line_spacing = spacing
        run = par.add_run(text)
        style_run(run, size=size, bold=bold)
        return par

    line('', 12, before=40)
    line('挑战杯·中国大学生创业计划竞赛', 14, bold=False, after=6)
    line(banner_line, 12, after=60)

    for t in titles:
        line(t, 26, bold=True, after=10, spacing=1.25)

    line('', 10, after=6)
    line(subtitle, 15, bold=False, before=6, after=50)

    for k, v in meta:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(6)
        par.paragraph_format.line_spacing = 1.4
        r1 = par.add_run(k + '：')
        style_run(r1, size=13, bold=True)
        r2 = par.add_run(v)
        style_run(r2, size=13)


def toc(doc, entries=None, line_spacing=1.18, l1_before=1, line_exact=None):
    """Table of contents.

    entries: list of (level, text, page). When provided, the TOC field is
    pre-populated with cached entries (so any viewer shows a real TOC) while
    remaining an updatable Word field. line_spacing / l1_before let a shorter
    document tighten the list; line_exact (in points) pins the line height so
    pagination does not depend on CJK font metrics.
    """
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(16)
    run = par.add_run('目    录')
    style_run(run, size=18, bold=True)

    if not entries:
        holder = doc.add_paragraph()
        add_field(holder, 'TOC \\o "1-2" \\h \\z \\u', '正在生成目录……')
        return

    instr = 'TOC \\o "1-2" \\h \\z \\u'
    last = len(entries) - 1
    for idx, (level, text, page) in enumerate(entries):
        par = doc.add_paragraph()
        pf = par.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(l1_before if level == 1 else 0)
        pf.line_spacing = Pt(line_exact) if line_exact else line_spacing
        para_mark_size(par, TABLE_SIZE)
        pf.tab_stops.add_tab_stop(Cm(USABLE_CM), WD_TAB_ALIGNMENT.RIGHT,
                                  WD_TAB_LEADER.DOTS)
        if level >= 2:
            indent_chars(par, left_chars=200)
        else:
            indent_chars(par, left_chars=0)

        if idx == 0:
            r0 = par.add_run()
            f1 = OxmlElement('w:fldChar')
            f1.set(qn('w:fldCharType'), 'begin')
            r0._r.append(f1)
            r1 = par.add_run()
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = instr
            r1._r.append(it)
            r2 = par.add_run()
            f2 = OxmlElement('w:fldChar')
            f2.set(qn('w:fldCharType'), 'separate')
            r2._r.append(f2)

        rt = par.add_run(text)
        style_run(rt, size=TABLE_SIZE, bold=(level == 1))
        rtab = par.add_run()
        rtab._r.append(OxmlElement('w:tab'))
        rp = par.add_run(str(page))
        style_run(rp, size=TABLE_SIZE, bold=(level == 1))

        if idx == last:
            r3 = par.add_run()
            f3 = OxmlElement('w:fldChar')
            f3.set(qn('w:fldCharType'), 'end')
            r3._r.append(f3)


def enable_update_fields(doc):
    settings = doc.settings.element
    el = OxmlElement('w:updateFields')
    el.set(qn('w:val'), 'true')
    settings.append(el)


def set_page_numbering(section, start=1, fmt=None):
    sectPr = section._sectPr
    pg = sectPr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType')
        sectPr.append(pg)
    pg.set(qn('w:start'), str(start))
    if fmt:
        pg.set(qn('w:fmt'), fmt)


def footer_page_number(section, prefix=''):
    section.footer.is_linked_to_previous = False
    par = section.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(0)
    if prefix:
        r = par.add_run(prefix)
        style_run(r, size=9)
    r = add_field(par, 'PAGE \\* arabic \\* MERGEFORMAT', '1')
    style_run(r, size=9)


def clear_footer(section):
    section.footer.is_linked_to_previous = False
    for par in section.footer.paragraphs:
        for r in list(par.runs):
            r._element.getparent().remove(r._element)


def header_text(section, text):
    section.header.is_linked_to_previous = False
    par = section.header.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(0)
    run = par.add_run(text)
    style_run(run, size=9, color='595959')
    para_border(par, edges=('bottom',), sz=6, color='BFBFBF')


def clear_header(section):
    section.header.is_linked_to_previous = False
    for par in section.header.paragraphs:
        for r in list(par.runs):
            r._element.getparent().remove(r._element)


def render(doc, ops):
    """Render a list of (kind, payload) operations."""
    for kind, payload in ops:
        if kind == 'h1':
            h1(doc, payload)
        elif kind == 'h2':
            h2(doc, payload)
        elif kind == 'h3':
            h3(doc, payload)
        elif kind == 'p':
            p(doc, payload)
        elif kind == 'p0':
            p(doc, payload, indent=False)
        elif kind == 'b1':
            bullet(doc, payload, 1)
        elif kind == 'b2':
            bullet(doc, payload, 2)
        elif kind == 'note':
            note(doc, payload[1], label=payload[0])
        elif kind == 'fig':
            fig(doc, payload)
        elif kind == 'cap':
            caption(doc, payload)
        elif kind == 'tbl':
            table(doc,
                  cols=payload['cols'],
                  rows=payload['rows'],
                  widths=payload.get('widths'),
                  caption_text=payload.get('caption'),
                  first_col_bold=payload.get('first_col_bold', False))
        elif kind == 'pb':
            page_break(doc)
        else:
            raise ValueError('unknown op: %s' % kind)
