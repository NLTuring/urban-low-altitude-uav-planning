# -*- coding: utf-8 -*-
"""Assemble the project brief DOCX (two-pass, with a real TOC)."""

import json
import os
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from docx import Document
from docx.enum.section import WD_SECTION

import docgen as G
from fontfix import fix_fonts_in_docx
from content_a import OPS_A
from content_b import OPS_B

OUT_DIR = os.path.dirname(HERE)
NAME = '城市低空多机协同调度与航路规划决策平台-项目简介与BP-PPT编制指南'
PDF_DIR = os.path.join(HERE, 'pdf')
TOC_JSON = os.path.join(HERE, 'toc.json')

_state = {'name': NAME}
OUT = os.path.join(OUT_DIR, NAME + '.docx')

DOC_TITLE = '城市低空多机协同调度与航路规划决策平台'
SOFFICE = r'C:\Program Files\LibreOffice\program\soffice.exe'

ALL_OPS = OPS_A + OPS_B
EXPECTED = [(1 if k == 'h1' else 2, v) for k, v in ALL_OPS if k in ('h1', 'h2')]


def build(entries):
    doc = G.new_document()

    G.cover(
        doc,
        titles=['城市低空多机协同调度', '与航路规划决策平台'],
        meta=[
            ('项目名称', '城市低空多机协同调度与航路规划决策平台'),
            ('参赛赛道', '挑战杯·中国大学生创业计划竞赛'),
            ('团队名称', '【待填】'),
            ('所在学校', '【待填】'),
            ('项目负责人', '【待填】'),
            ('指导教师', '【待填】'),
            ('联系方式', '【待填】'),
        ],
    )
    cover_sec = doc.sections[0]
    G.clear_header(cover_sec)
    G.clear_footer(cover_sec)

    doc.add_section(WD_SECTION.NEW_PAGE)
    toc_sec = doc.sections[1]
    G.clear_header(toc_sec)
    G.clear_footer(toc_sec)
    G.toc(doc, entries, line_exact=14)

    doc.add_section(WD_SECTION.NEW_PAGE)
    body_sec = doc.sections[2]
    G.header_text(body_sec, DOC_TITLE)
    G.footer_page_number(body_sec)
    G.set_page_numbering(body_sec, start=1, fmt='decimal')

    G.render(doc, ALL_OPS)
    G.enable_update_fields(doc)

    props = doc.core_properties
    props.title = DOC_TITLE + ' 项目简介暨商业计划书与路演PPT编制指南'
    props.subject = '挑战杯中国大学生创业计划竞赛 参赛项目简介'
    props.author = '【待填：团队名称】'

    name = _state['name']
    out_path = os.path.join(OUT_DIR, name + '.docx')
    try:
        doc.save(out_path)
    except PermissionError:
        if name == NAME:
            name = NAME + '(新版本)'
            _state['name'] = name
            print('!! 原文档被占用（可能已在 WPS / Word 中打开），改存为 %s.docx' % name)
        out_path = os.path.join(OUT_DIR, name + '.docx')
        doc.save(out_path)
    fix_fonts_in_docx(out_path)
    return doc


def to_pdf(name):
    os.makedirs(PDF_DIR, exist_ok=True)
    subprocess.run([SOFFICE, '--headless', '--convert-to', 'pdf',
                    '--outdir', PDF_DIR, os.path.join(OUT_DIR, name + '.docx')],
                   check=True, timeout=900)
    return os.path.join(PDF_DIR, name + '.pdf')


def main():
    # ---- pass 1: build without a populated TOC, then read real page numbers ----
    build(None)
    pdf1 = to_pdf(_state['name'])
    subprocess.run([sys.executable, os.path.join(HERE, 'extract_toc.py'),
                    pdf1, TOC_JSON], check=True, cwd=HERE)

    rows = json.load(open(TOC_JSON, encoding='utf-8'))
    if len(rows) != len(EXPECTED):
        print('!! heading count mismatch: extracted=%d expected=%d'
              % (len(rows), len(EXPECTED)))
        print('   extracted text:', [r['text'] for r in rows])
    else:
        print('heading count matches: %d entries' % len(rows))

    entries = []
    for i, (level, text) in enumerate(EXPECTED):
        page = rows[i]['page'] if i < len(rows) else 0
        entries.append((level, text.replace('\u3000', ' ').strip(), page))

    # ---- pass 2: rebuild with the populated TOC ----
    name = _state['name']
    build(entries)
    name = _state['name']
    pdf2 = to_pdf(name)
    print('final pdf:', pdf2)

    import shutil
    preview = os.path.join(OUT_DIR, name + '(预览版).pdf')
    try:
        shutil.copyfile(pdf2, preview)
        print('preview pdf:', preview)
    except PermissionError:
        print('!! 预览 PDF 被占用，未更新')

    import pymupdf
    d = pymupdf.open(pdf2)
    print('pages:', d.page_count)
    doc = Document(os.path.join(OUT_DIR, name + '.docx'))
    print('docx:', os.path.join(OUT_DIR, name + '.docx'))
    print('paragraphs=%d tables=%d sections=%d'
          % (len(doc.paragraphs), len(doc.tables), len(doc.sections)))


if __name__ == '__main__':
    main()
